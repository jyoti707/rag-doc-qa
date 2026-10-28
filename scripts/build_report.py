import os, json, pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
FIG = os.path.join(ROOT, "figures")
ev = json.load(open(os.path.join(ROOT, "data", "eval_summary_hashing.json")))
cl = json.load(open(os.path.join(ROOT, "data", "cleaning_report.json")))
ab = pd.read_csv(os.path.join(ROOT, "data", "ablation_results.csv"))
rf = pd.read_csv(os.path.join(ROOT, "data", "refusal_test.csv"))
doc = Document()
sec = doc.sections[0]; sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for m in ("left_margin", "right_margin"): setattr(sec, m, Cm(2.2))
sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(2.0)
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
for n, sz, col in [("Heading 1", 16, "1F3864"), ("Heading 2", 13, "2F5597"), ("Heading 3", 11, "2F5597")]:
    h = doc.styles[n]; h.font.name = "Calibri"; h.font.size = Pt(sz); h.font.bold = True; h.font.color.rgb = RGBColor.from_string(col)

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), color); tcPr.append(s)

def P(text="", bold=False, italic=False, size=None, align=None, color=None, after=6):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(after); return p

def RP(parts, style=None):
    """paragraph with mixed bold: parts = [(text, bold), ...]"""
    p = doc.add_paragraph(style=style)
    for t, b in parts: r = p.add_run(t); r.bold = b
    p.paragraph_format.space_after = Pt(4); return p

def B(text, lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if lead: r = p.add_run(lead); r.bold = True
    p.add_run(text); p.paragraph_format.space_after = Pt(2); return p

def N(text, lead=None):
    p = doc.add_paragraph(style="List Number")
    if lead: r = p.add_run(lead); r.bold = True
    p.add_run(text); p.paragraph_format.space_after = Pt(2); return p

def T(rows, widths=None, header=True, fs=9):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    if widths:
        for j, w in enumerate(widths): t.columns[j].width = Cm(w)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j); c.text = ""; p = c.paragraphs[0]; r = p.add_run(str(val)); r.font.size = Pt(fs)
            if header and i == 0: r.bold = True; shade(c, "D9E2F3")
            if widths: c.width = Cm(widths[j])
    doc.add_paragraph().paragraph_format.space_after = Pt(2); return t

def F(name, caption, w=16):
    doc.add_picture(os.path.join(FIG, name), width=Cm(w)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    P(caption, italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

H1 = lambda t: doc.add_heading(t, 1); H2 = lambda t: doc.add_heading(t, 2); H3 = lambda t: doc.add_heading(t, 3)

# footer page numbers
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("RAG-Based Document Q&A System  |  HCLTech Career Shaper Programme 2025-26  |  Page "); r.font.size = Pt(8)
for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
    run = fp.add_run(); run.font.size = Pt(8)
    if tag: e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag); run._r.append(e)
    else: e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; run._r.append(e)

# ============================ TITLE
P("Project Report & Synopsis", bold=True, size=24, align=WD_ALIGN_PARAGRAPH.CENTER, color="1F3864", after=2)
P("RAG-Based Document Q&A System", bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER, color="2F5597", after=2)
P("HCLTech Career Shaper Programme | III Year Industry Project | 2025-26", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
T([["Field", "Details (Learner to complete)"], ["Project Title", "RAG-Based Document Q&A System"], ["Project No.", "72"], ["Domain", "LLM / Python"],
   ["Skill / Tech Stack", "Python, LangChain (text splitters), ChromaDB, HuggingFace sentence-transformers, OpenAI-compatible APIs, RAG, Vector DB, Prompt Engineering, Streamlit"],
   ["Difficulty", "Advanced"], ["Duration / Format", "4 Weeks | Individual Project"], ["Learner Name", "[FILL IN]"], ["Roll Number", "[FILL IN]"],
   ["College / Campus", "[FILL IN]"], ["Mentor Name", "[FILL IN]"], ["Submission Date", "[FILL IN]"]], widths=[4.5, 12])

# ============================ 1
H1("1. Problem Statement")
H2("1.1 Background")
P("Organisations and students keep knowledge in long PDF, Word, HTML and text documents: rules, policies, manuals and scheme guidelines. Finding one fact usually means opening several files and reading pages. Keyword search only finds exact words, and a general-purpose chatbot does not know the private documents and may invent an answer that sounds correct (a hallucination).")
P("Retrieval-Augmented Generation (RAG) solves this by splitting the documents into small chunks, storing their embeddings in a vector database, retrieving the chunks most similar to the question and asking a Large Language Model (LLM) to answer only from those chunks. The user can then verify the answer against the displayed source text. This project builds such a system end to end, measures how well it works, and is careful about its limits.")
H2("1.2 Problem statement (5W1H)")
T([["Question", "Answer"],
   ["Who faces the problem?", "Students, employees, support teams and researchers who must find answers inside many long documents (for example constitutional articles, government scheme rules or technical manuals)."],
   ["What is the problem?", "Finding a specific, correct answer in large document collections is slow, and ungrounded chatbots can hallucinate without showing where the answer came from."],
   ["When / where does it occur?", "Whenever a person needs a quick, reliable answer from internal or public documents: exam preparation, helpdesks, onboarding, compliance checks."],
   ["Why is it important?", "Wrong answers about rules, entitlements or procedures cause real harm. Trust needs a visible source and a system that says \"not found\" instead of guessing."],
   ["How will it be solved?", "A Streamlit web application that indexes PDF/Word/HTML/Markdown/TXT files into ChromaDB, retrieves the top-k chunks for each question, generates a cited answer with an LLM (or offline extractive mode), shows the source chunks, scores retrieval and grounding, and refuses when evidence is weak."]], widths=[3.8, 12.7])
H2("1.3 Objectives")
for t in ["Ingest documents in five formats (PDF, Word, HTML, Markdown, TXT) and split them into overlapping chunks with LangChain's text splitter.",
          "Create embeddings (MiniLM, OpenAI or an offline hashing embedder) and store them in a persistent ChromaDB vector store.",
          "Retrieve the top-k most similar chunks and generate a short answer with citations using an OpenAI-compatible or local open-source LLM.",
          "Always display the source chunks next to the answer, with retrieval score and a grounding score, as the ethical note of the brief requires.",
          "Refuse to answer (and say so) when the best retrieval score is below a threshold, to reduce hallucination.",
          "Build a corpus of 50 public-domain documents with 1,000 Q&A pairs and measure hit rate, MRR and answer quality.",
          "Expand and clean dataset.csv (500 to 5,000 rows), document it in data_dictionary.csv and analyse it in an EDA notebook with 3+ visualisations.",
          "Work at zero cost: offline mode, free API tiers (Groq, Gemini) and local Ollama are supported."]: B(t)
H2("1.4 Scope and limitations")
RP([("In scope: ", True), ("English text documents, question answering with sources, vector search, query log and feedback, retrieval evaluation, dataset EDA.", False)])
RP([("Out of scope / limitations:", True)])
for t in ["Scanned (image-only) PDFs are not read, because there is no OCR.", "Tables, figures and diagrams inside documents are not understood; only text is indexed.",
          "The built-in offline hashing embedder matches words, not meaning; semantic questions need the MiniLM or OpenAI embedder.",
          "LLM answers can still be wrong. The tool gives a draft answer with evidence, not a final authority. It must not be used for legal or medical advice without expert review.",
          "The supplied dataset.csv is synthetic (see section 5); conclusions about real behaviour come from the real corpus experiments."]: B(t)

# ============================ 2
H1("2. Proposed Solution & Approach")
H2("2.1 Solution overview")
P("The application is called the RAG-Based Document Q&A System. The central design idea is that the program, not the LLM, controls what the model may see and say. The program decides which chunks are retrieved, builds a strict prompt (\"answer only from the numbered passages, otherwise say you could not find it\"), checks the answer against the passages and decides whether to show it or refuse. The LLM only writes the final sentence from evidence that the user can also read.")
H2("2.2 How the system works")
for lead, t in [("Load. ", "A file is read with pypdf (PDF), python-docx (Word), a small HTML parser, or plain text reading (MD/TXT)."),
                ("Chunk. ", "LangChain's RecursiveCharacterTextSplitter cuts the text into chunks of about 500 characters with 80 characters of overlap, preferring paragraph and sentence boundaries."),
                ("Embed. ", "Each chunk is converted into a vector by the chosen embedder (all-MiniLM-L6-v2, OpenAI text-embedding-3-small or the offline hashing embedder)."),
                ("Store. ", "Vectors, chunk text and metadata (doc_id, title, doc_type, chunk number) are saved in ChromaDB with cosine distance, on disk, so the index survives restarts."),
                ("Retrieve. ", "The question is embedded and the top-k similar chunks are returned with their cosine similarity scores."),
                ("Gate. ", "If the best score is below the minimum retrieval score, the system answers \"I could not find this in the documents\" and still shows the nearest chunks."),
                ("Generate. ", "The chunks are placed in a numbered context and sent with the question to the LLM (OpenAI, Groq, Gemini, Ollama or Claude). In offline mode the best-matching sentences are extracted instead."),
                ("Guard and display. ", "A grounding score (share of answer words that appear in the sources) is computed. Low grounding triggers a warning. The answer, scores and source chunks are displayed, and the question is logged for the analytics tab."),
                ("Learn. ", "The user can mark the answer helpful or wrong; the feedback is stored in SQLite.")]: N(t, lead)
H2("2.3 Key features")
T([["Feature", "What it does", "Benefit"],
   ["Multi-format ingestion", "PDF, Word, HTML, Markdown, TXT upload or bundled sample corpus.", "Works with real enterprise documents."],
   ["Persistent vector store", "ChromaDB on disk with cosine similarity and metadata.", "No re-indexing on restart; fast search."],
   ["Source-first answers", "Every answer lists the retrieved chunks with document, type and score.", "User can verify; supports the ethical requirement."],
   ["Refusal threshold", "Refuses when retrieval is weak.", "Fewer made-up answers."],
   ["Grounding score", "Checks how much of the answer is supported by the context.", "Flags possible hallucination."],
   ["Multiple providers", "Offline extractive, OpenAI, Groq, Gemini, Ollama, Claude.", "Works with no money and no internet."],
   ["Query log and feedback", "SQLite log with scores and thumbs up/down; CSV download.", "Monitoring and improvement."],
   ["Built-in evaluation", "Scripts for hit@k, MRR, ablation and refusal tests on 1,000 Q&A pairs.", "Measurable performance."]], widths=[3.6, 7.2, 5.7])
H2("2.4 Key design decisions")
for lead, t in [("Own thin RAG layer over heavy frameworks. ", "Only LangChain's text splitter is used; retrieval, prompting and guarding are small readable functions. This keeps the code easy to explain, test and debug, and avoids breakage from frequent framework API changes."),
                ("Embedder abstraction. ", "One interface with three back-ends lets the system run offline for testing and switch to MiniLM or OpenAI for better meaning-based search."),
                ("Provider abstraction. ", "Gemini, Groq, Ollama and OpenAI share the OpenAI-compatible API, so one wrapper serves all; Claude uses its own SDK."),
                ("Refuse rather than guess. ", "A threshold, a strict prompt and a grounding check are combined, because a wrong confident answer is worse than no answer."),
                ("SQLite and ChromaDB on disk. ", "No server is required, data stays on the user's laptop.")]: B(t, lead)
H2("2.5 Development approach and 4-week plan")
T([["Week", "Activities", "Deliverables"],
   ["Week 1", "Problem study, tool setup, EDA on dataset.csv, corpus design, architecture and prompt design.", "Problem statement, EDA notebook, 50-document corpus"],
   ["Week 2", "Loaders, chunking, embeddings, ChromaDB store, retrieval and first answers (Sprint 1).", "Working retrieval module demo"],
   ["Week 3", "LLM providers, guards, Streamlit UI, logging, evaluation scripts, functional tests, ablation.", "End-to-end system and recorded metrics"],
   ["Week 4", "Report, screenshots, demo video or live demo, ethical declaration.", "All deliverables submitted"]], widths=[2.2, 9.5, 4.8])

# ============================ 3
H1("3. System Architecture")
P("The architecture is shown at three levels: the context of the system (DFD Level 0), its internal processes (DFD Level 1) and its software structure (block diagram).")
H2("3.1 Data Flow Diagram - Level 0"); F("arch_dfd0.png", "Figure 1: DFD Level 0. The user supplies documents and questions and receives answers with sources.", 15)
H2("3.2 Data Flow Diagram - Level 1"); F("arch_dfd1.png", "Figure 2: DFD Level 1. Indexing (top row) runs once per document; querying (bottom row) runs for every question.", 16)
T([["Process", "Input", "Processing", "Output"],
   ["1.0 Load", "Uploaded file", "Extract text by file type; clean blank lines", "Plain text"],
   ["2.0 Chunk", "Plain text", "Recursive split, 500 chars, 80 overlap", "Chunks"],
   ["3.0 Embed", "Chunks / question", "Embedding model", "Vectors"],
   ["4.0 Store", "Vectors, text, metadata", "Upsert into ChromaDB (cosine)", "Persistent index (D1)"],
   ["5.0-6.0 Retrieve", "Question vector", "Top-k similarity search", "Chunks with scores"],
   ["7.0 Generate", "Chunks, question", "Prompt + LLM or extractive answer", "Draft answer"],
   ["8.0 Guard", "Draft answer, scores", "Threshold, grounding score, warning", "Final answer, log row (D2)"]], widths=[2.8, 3.5, 6.0, 4.2])
H2("3.3 System block diagram"); F("arch_blocks.png", "Figure 3: Three-layer structure: presentation (Streamlit), core logic (package rag/) and external/storage layer.", 15)
H2("3.4 Module description")
T([["Module", "Responsibility", "Important functions"],
   ["app.py", "Streamlit UI: settings, document tab, ask tab, analytics tab, evaluation tab.", "get_pipe, tab handlers"],
   ["rag/loaders.py", "Read PDF, DOCX, HTML, MD, TXT.", "load_file"],
   ["rag/embeddings.py", "Embedding back-ends and default thresholds.", "get_embedder, HashingEmbedder, MiniLMEmbedder, OpenAIEmbedder"],
   ["rag/store.py", "ChromaDB wrapper.", "add, query, count, doc_summary, reset"],
   ["rag/pipeline.py", "Chunking, indexing, retrieval, answering, guards.", "split_text, index_file, retrieve, ask, grounding_score"],
   ["rag/llm.py", "All LLM providers, prompt, retry on rate limits.", "LLM.complete, PROVIDERS, SYSTEM_PROMPT"],
   ["rag/logdb.py", "Query log and feedback in SQLite.", "log, set_feedback, read"],
   ["scripts/", "Corpus building, dataset expansion and cleaning, evaluation, ablation, refusal test, diagrams, notebook and report builders.", "build_corpus, expand_dataset, clean_dataset, evaluate"]], widths=[3.2, 7.3, 6.0])
H2("3.5 Key algorithms")
RP([("(a) Chunking. ", True), ("Recursive splitting tries separators in the order paragraph, line, sentence, space, so chunks end at natural boundaries. Overlap repeats the last 80 characters in the next chunk so that a fact at a boundary is not lost.", False)])
RP([("(b) Retrieval. ", True), ("Cosine similarity between the question vector and chunk vectors. ChromaDB returns cosine distance d, and the score is 1 - d. Top-k chunks are returned in decreasing score.", False)])
RP([("(c) Offline hashing embedder. ", True), ("Words and word pairs are hashed into a 4,096-dimension vector (scikit-learn HashingVectorizer, stop-words removed, L2-normalised). It needs no download and no fitting, so new documents can be added at any time; it matches vocabulary rather than meaning.", False)])
RP([("(d) Refusal gate. ", True), ("If the top score is below the threshold (0.15 for hashing, 0.30 for MiniLM/OpenAI), no answer is generated.", False)])
RP([("(e) Grounding score. ", True), ("Content words (stop-words and very short words removed) of the answer are compared with the words of the retrieved context: grounding = |answer words in context| / |answer words|. Below 0.6 the app warns the user.", False)])
RP([("(f) Extractive answer. ", True), ("In offline mode the sentences of the retrieved chunks are ranked by overlap with the question plus chunk score, and the best two are returned with citation numbers.", False)])
H2("3.6 Data design")
T([["Store", "Content"],
   ["ChromaDB collection rag_<embedder>", "id = DOCxxx_Cnnn; document = chunk text; metadata = doc_id, title, doc_type, chunk_no; embedding vector; cosine index."],
   ["qa_log.db, table query_log", "id, ts, question, answer, top_score, grounding, abstained, provider, embedder, sources, feedback."],
   ["data/qa_pairs.csv", "qa_id, doc_id, doc_title, source, query_type, question, answer, evidence (1,000 rows)."],
   ["data/dataset.csv", "doc_id, doc_type, chunk_id, word_count, retrieval_score, query_type, answer_quality_score (5,000 rows)."]], widths=[5.2, 11.3])

# ============================ 4
H1("4. Technology Stack")
T([["Layer", "Technology used"],
   ["Frontend / UI", "Streamlit (tabs, file uploader, metrics, expanders, charts)."],
   ["Backend / Logic", "Python 3.10+; package rag/; langchain-text-splitters; re, dataclasses, sqlite3."],
   ["Vector database", "ChromaDB (persistent client, cosine space)."],
   ["Embeddings", "sentence-transformers all-MiniLM-L6-v2 (384-d), OpenAI text-embedding-3-small, or offline HashingVectorizer."],
   ["LLM", "OpenAI-compatible API: OpenAI, Groq, Gemini, Ollama (local); Anthropic Claude via its SDK; offline extractive mode."],
   ["Document parsing", "pypdf, python-docx, html.parser."],
   ["Analysis / visualisation", "pandas, numpy, scipy, matplotlib, seaborn, Jupyter."],
   ["Hardware", "Laptop with 8 GB RAM, no GPU needed (cloud LLM or extractive mode)."]], widths=[4, 12.5])
H2("4.1 Why these choices")
for lead, t in [("ChromaDB ", "is serverless, open source and pairs with Python; the brief specifies it."),
                ("MiniLM ", "is small (about 90 MB), fast on CPU and good for English semantic search."),
                ("Streamlit ", "lets one developer build a complete interface in Python and run it locally."),
                ("Free providers ", "(Groq, Gemini) and Ollama keep the cost at zero; OpenAI remains supported because the brief lists it.")]: B(t, lead)

# ============================ 5
H1("5. Dataset Used")
T([["Template field", "Details"],
   ["Dataset file name", "data/dataset.csv (cleaned). Original: data/dataset_original_500.csv. Corpus: data/corpus/ (50 files). Q&A: data/qa_pairs.csv."],
   ["Number of rows", "5,000 in dataset.csv (500 original + 4,500 synthetic); 1,000 Q&A pairs; 50 documents (169 chunks at 500 characters)."],
   ["Number of columns", "7 in dataset.csv; 13 in dataset_features.csv (6 derived columns)."],
   ["Key columns", "doc_id, doc_type, chunk_id, word_count, retrieval_score, query_type, answer_quality_score."],
   ["Additional data collected", "Optional query log (qa_log.db) written locally when the app is used; no personal data."]], widths=[4.2, 12.3])
H2("5.1 Corpus of 50 documents and 1,000 Q&A pairs")
P("As requested in the dataset-generation prompt, a corpus of 50 public-information documents was created: 18 on the Indian Constitution (Preamble, Fundamental Rights, Directive Principles, President, Parliament, courts, amendments, emergency, federalism, elections, schedules), 16 on government schemes (PM-KISAN, Ayushman Bharat, Jan Dhan, Ujjwala, Swachh Bharat, Digital India, Mudra and others) and 16 technical manuals (Python, Git, Linux, SQL, HTTP, Docker, networking, RAG, vector databases, cybersecurity and others). The documents were written in the author's own words from public facts, saved as 10 TXT, 10 Markdown, 10 HTML, 10 Word and 10 PDF files, so every loader is exercised.")
P("Each document holds 8 facts, and every fact has a question and a short answer. The 1,000 Q&A pairs consist of 400 original questions, 400 variants that mention the document title, 150 variants with an instruction prefix and 50 summarisation questions. The variants test robustness to wording but are not independent questions; this is a limitation (section 7). The scripts/build_corpus.py file regenerates everything.")
T([["Query type", "Count", "Example"], ["factual", "686", "When did the Constitution of India come into force?"], ["analytical", "161", "What does Article 21 protect?"],
   ["procedural", "84", "Which command creates a virtual environment?"], ["summarization", "50", "Give a short summary of the document on Docker Basics."], ["comparative", "19", "What is the difference between a list and a tuple?"]], widths=[3.5, 2.2, 10.8])
H2("5.2 Dataset expansion and cleaning")
P("The supplied dataset.csv has 500 rows. Following the generation prompt (4,500 additional rows to reach 5,000), scripts/expand_dataset.py created new rows with the same 7 columns, resampling each column from the original ranges and class proportions, continuing chunk_id from CHK00501 and keeping doc_type consistent with the real file format of each document. These rows are synthetic and are labelled so in the report. The prompt's instruction to label anomalies at 2-5% was not applicable because the file has no timestamp or anomaly column; instead a hallucination-risk flag was derived (section 5.3).")
T([["Cleaning step", "Result"], ["Rows in / out", f"{cl['rows_in']} / {cl['rows_out']}"], ["Missing values", str(cl["missing_values"])], ["Duplicate rows / duplicate chunk_id", f"{cl['duplicate_rows']} / {cl['duplicate_chunk_ids']}"],
   ["doc_ids with more than one doc_type (before)", f"{cl['doc_ids_with_multiple_types_before']} of 50 (impossible for one real file)"], ["doc_type values corrected to match corpus", str(cl["doc_type_rows_corrected"])],
   ["Out-of-range rows (scores outside 0-1)", str(cl["out_of_range_rows_dropped"])], ["Derived columns added", "length_bucket, retrieval_band, quality_band, combined_score, hallucination_risk, category"]], widths=[7, 9.5])
H2("5.3 Data dictionary")
P("All columns are described in data_dictionary.csv (type, sample values, nullability, unique count, range and a written description). Summary of the original columns: doc_id (identifier of one of 50 documents), doc_type (file format), chunk_id (unique chunk), word_count (words in the chunk, 50-499), retrieval_score (cosine similarity, 0.50-1.00), query_type (five classes) and answer_quality_score (0.60-1.00).")
H2("5.4 Exploratory data analysis")
F("fig1_distributions.png", "Figure 4: Distributions of word_count, retrieval_score and answer_quality_score (nearly uniform).", 16)
F("fig2_categories.png", "Figure 5: Number of records per document type, query type and corpus category.", 16)
F("fig4_correlation.png", "Figure 6: Correlation matrix and retrieval-vs-quality density; correlations are about zero.", 15)
F("fig5_risk.png", "Figure 7: Share of records meeting the hallucination-risk rule, by category and format.", 15)
P("Findings. (1) The three numeric columns are almost uniformly distributed with no outliers. (2) Retrieval score, answer quality and chunk length are uncorrelated (|r| < 0.02), and Kruskal-Wallis tests show no significant difference between question types (p = 0.72 for answer quality, p = 0.59 for retrieval score across formats). (3) About 29.6% of records meet the rule retrieval_score < 0.6 or answer_quality_score < 0.65 and would need human review. (4) The lack of any relationship shows that the dataset is synthetic; a real RAG log would normally show higher answer quality for higher retrieval scores. Section 6 therefore also reports measurements on the real corpus.")

# ============================ 6
H1("6. Key Results & Metrics")
P("Part A reports functional tests; Part B reports quality and performance measured on the real 50-document corpus with the offline hashing embedder and extractive answers (no internet, no API key). Part C lists the measurements that must be repeated with MiniLM and a real LLM on the learner's laptop.")
H2("6.1 Part A - Verified functional results")
T([["Test (tests/test_rag.py)", "What was checked", "Result"],
   ["Loaders, all formats", "TXT, MD, HTML, DOCX and PDF are read and produce clean text.", "Pass"], ["HTML script removal", "Script content is not indexed.", "Pass"],
   ["Unsupported file type", "Clear error for unknown extension.", "Pass"], ["Chunking", "Chunk size limit respected; neighbouring chunks overlap.", "Pass"],
   ["Grounding score", "Supported answer scores > 0.9; invented answer scores < 0.3.", "Pass"], ["Index, retrieve, answer", "Constitution question returns 1950 with DOC001 as top source.", "Pass"],
   ["Out-of-domain refusal", "Bitcoin question is refused; sources still shown.", "Pass"], ["Empty index", "Friendly message instead of crash.", "Pass"],
   ["Metadata", "Each chunk carries doc_id, title, doc_type, chunk_no, score.", "Pass"], ["Query log and feedback", "Row written and thumbs-up stored.", "Pass"],
   ["Re-indexing", "Indexing the same file twice does not duplicate chunks.", "Pass"], ["Total", "", "11 of 11 passed"]], widths=[4.3, 9.2, 3])
P("The Streamlit app was also exercised with Streamlit's AppTest harness: index the sample corpus, ask an in-domain question, then an out-of-domain question, without exceptions.")
H2("6.2 Part B - Retrieval and answer quality on the real corpus")
T([["Metric", "How measured", "Result"],
   ["Document hit@4", "Correct document among top-4 chunks, 1,000 questions", f"{ev['doc_hit_at_k']*100:.1f}%"],
   ["Top-1 accuracy", "Correct document is the first result", f"{ev['top1_accuracy']*100:.1f}%"],
   ["Mean Reciprocal Rank (MRR)", "Average of 1/rank of correct document", f"{ev['mrr']:.3f}"],
   ["Evidence recall@4", "Chunk containing the answer sentence retrieved", f"{ev['evidence_recall_at_k']*100:.1f}%"],
   ["Average top retrieval score", "Mean cosine similarity of best chunk", f"{ev['avg_top_score']:.2f}"],
   ["Average grounding score", "Extractive answers vs their context", f"{ev['avg_grounding']:.2f}"],
   ["Answer word recall", "Gold-answer words found in the generated answer", f"{ev['avg_answer_recall']*100:.1f}%"],
   ["Abstain rate", "Questions refused at threshold 0.15", f"{ev['abstain_rate']*100:.1f}%"],
   ["Time per query", "Embedding + search + extractive answer, CPU", f"{ev['seconds_per_query']*1000:.0f} ms"]], widths=[4.5, 8, 4])
P("Important reading: the questions were written from the same facts as the documents, and 40% of them carry the title or similar wording, so these numbers are optimistic. On the 400 'hard' questions that do not mention the title, document hit@1 is 90.0% and hit@5 is 99.0% (chunk size 500).")
H3("Ablation: chunk size and top-k (questions without title, n = 400)")
hard = ab[(ab.subset == "hard_no_title")]
rows = [["Chunk size / overlap", "Chunks", "k", "Doc hit", "MRR", "Evidence recall"]]
for _, r in hard.iterrows(): rows.append([f"{int(r.chunk_size)} / {int(r.overlap)}", int(r.chunks), int(r.k), f"{r.doc_hit:.3f}", f"{r.mrr:.3f}", f"{r.evidence_recall:.3f}"])
T(rows, widths=[3.8, 2, 1.5, 2.8, 2.8, 3.6], fs=8.5)
P("Larger chunks raise evidence recall (0.835 at k=1 for 800 characters versus 0.688 for 300) because a fact is less likely to be split; document hit rate hardly changes. Raising k from 3 to 5 gives a small gain at the price of more context for the LLM.")
F("fig8_real_retrieval.png", "Figure 8: Real retrieval experiment: hit rate versus top-k, evidence recall, and evidence found by retrieval-score band.", 16.5)
H3("Refusal threshold test")
rows = [["Threshold", "Out-of-domain questions refused (n=20)", "In-domain questions wrongly refused (n=200)"]]
for _, r in rf.iterrows(): rows.append([f"{r.threshold:.2f}", f"{r.ood_refused*100:.0f}%", f"{r.in_domain_wrongly_refused*100:.1f}%"])
T(rows, widths=[3, 6.5, 7])
r15 = rf[rf.threshold == 0.15].iloc[0]
P(f"A threshold of 0.15 refuses {r15.ood_refused*100:.0f}% of unrelated questions while wrongly refusing {r15.in_domain_wrongly_refused*100:.1f}% of valid ones for the hashing embedder; this is the default. (Refusal also happens when no retrieved sentence shares a word with the question.) Semantic embedders separate the two groups better and use 0.30.")
H2("6.3 Part C - To be measured on your laptop with MiniLM and a real LLM")
P("Run: python scripts/evaluate.py --embedder minilm --provider groq --n 100 --reindex (set the API key first). Record the values in the table below and add one teacher rating.")
T([["Metric", "How measured", "Result"],
   ["Doc hit@4 and MRR with MiniLM", "evaluate.py --embedder minilm", "[FILL IN]"], ["Answer word recall with LLM", "evaluate.py --provider ...", "[FILL IN]"],
   ["Average grounding with LLM", "evaluate.py --provider ...", "[FILL IN]"], ["Abstain rate with LLM", "evaluate.py --provider ...", "[FILL IN]"],
   ["Time per answer with LLM", "Stopwatch / seconds_per_query", "[FILL IN: seconds]"], ["Teacher / mentor rating (1-5)", "Rate 20 answers for correctness and clarity", "[FILL IN]"]], widths=[5, 7, 4.5])
H2("6.4 Business value achieved")
for lead, t in [("Faster answers. ", "A question is answered in milliseconds (offline) to a few seconds (LLM) instead of minutes of manual searching."),
                ("Trust. ", "Source chunks, scores, grounding check and refusal make errors visible."),
                ("Low cost and privacy. ", "Offline mode and Ollama keep data on the laptop; free tiers cover cloud use."),
                ("Reusable assets. ", "Corpus, Q&A pairs, evaluation scripts and the query log can be reused for any new document set.")]: B(t, lead)

# ============================ 7
H1("7. Challenges Faced & How Resolved")
T([["#", "Challenge", "Resolution"],
   ["1", "Hallucination: a language model may answer from memory when retrieval is poor.", "Strict system prompt, refusal threshold, grounding score and always-visible source chunks."],
   ["2", "No internet or model download in the test environment.", "Added an offline hashing embedder and extractive answer mode, so the whole pipeline is testable without any key."],
   ["3", "Facts cut at chunk boundaries lowered evidence recall (0.69 at 300 characters).", "Overlap, recursive splitting and a larger default option; ablation documents the effect."],
   ["4", "Original dataset.csv had several doc_type values for the same doc_id.", "Aligned doc_type with the real file format in the corpus manifest (390 rows corrected)."],
   ["5", "Dataset columns were uncorrelated, so EDA could not show real behaviour.", "Declared the data synthetic and added a real retrieval experiment on 1,000 Q&A pairs."],
   ["6", "Q&A pairs share wording with documents, so scores are optimistic.", "Reported a harder subset without titles; recommend generating extra questions with an LLM (scripts/generate_qa_with_llm.py) and testing with MiniLM."],
   ["7", "Rate limits of free LLM tiers (HTTP 429).", "Automatic retry with waiting, small context, and alternative providers."],
   ["8", "Different file formats (PDF, Word, HTML) give messy text.", "Per-format loaders and whitespace cleaning; scanned PDFs documented as a limitation."],
   ["9", "Setup errors on a fresh laptop (missing packages, wrong folder).", "Fixed requirements.txt, virtual environment guide and a troubleshooting table in README."]], widths=[0.8, 7.2, 8.5])

# ============================ 8
H1("8. References")
refs = ["Lewis, P., et al. (2020). Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks. Advances in Neural Information Processing Systems 33 (NeurIPS 2020).",
        "Vaswani, A., et al. (2017). Attention Is All You Need. Advances in Neural Information Processing Systems 30.",
        "Brown, T. B., et al. (2020). Language Models are Few-Shot Learners. Advances in Neural Information Processing Systems 33.",
        "Reimers, N., & Gurevych, I. (2019). Sentence-BERT: Sentence Embeddings using Siamese BERT-Networks. EMNLP-IJCNLP 2019.",
        "Karpukhin, V., et al. (2020). Dense Passage Retrieval for Open-Domain Question Answering. EMNLP 2020.",
        "Ji, Z., et al. (2023). Survey of Hallucination in Natural Language Generation. ACM Computing Surveys, 55(12).",
        "Chroma. ChromaDB documentation. https://docs.trychroma.com",
        "LangChain. Text splitters documentation. https://python.langchain.com",
        "Streamlit Inc. Streamlit documentation. https://docs.streamlit.io",
        "Hugging Face. sentence-transformers/all-MiniLM-L6-v2 model card. https://huggingface.co/sentence-transformers/all-MiniLM-L6-v2",
        "OpenAI, Groq, Google and Anthropic API documentation (provider pages).",
        "Government of India. The Constitution of India. https://legislative.gov.in",
        "Government of India. The Digital Personal Data Protection Act, 2023 (Act No. 22 of 2023).",
        "HCLTech Career Shaper Programme Guidelines (provided by Programme Coordinator). [Mentor to validate and add links.]"]
for i, r in enumerate(refs, 1): RP([(f"{i}. ", True), (r, False)])

# ============================ 9
H1("9. Ethical Compliance Declaration")
P("I confirm that:")
for t in ["☐  All dataset usage complies with the source licence.", "☐  No personally identifiable information (PII) is stored without consent.",
          "☐  AI tools used for code/writing are disclosed in the Appendix.", "☐  This project complies with the DPDP Act 2023 where applicable.",
          "☐  I understand that RAG systems may hallucinate, that the source chunk is always displayed next to the answer, and that the system is not used for legal or medical advice without expert review."]: P(t, after=2)
H2("9.1 How each point is addressed")
for lead, t in [("Licence. ", "The corpus was written by the author in their own words from public facts; no third-party text is copied. Libraries are open source (Streamlit, ChromaDB, LangChain splitters, pypdf, python-docx, scikit-learn, pandas, matplotlib) and used under their licences; each AI provider is used under its terms. The Indian Constitution is a public document of the Government of India."),
                ("Personal data. ", "No names, contacts or student data are collected. API keys are typed in the sidebar or read from .env, stay in the session and are never logged. The query log is a local file."),
                ("Data sent to providers. ", "When a cloud LLM is chosen, the question and retrieved chunks are sent to that provider. Private documents should use extractive mode or local Ollama."),
                ("Hallucination and responsible use. ", "Source chunks are always shown; weak retrieval is refused; low grounding triggers a warning; the app and README state that the output is a draft and not legal or medical advice."),
                ("Dataset honesty. ", "4,500 of 5,000 rows in dataset.csv are synthetic and are described as such; real findings come from the corpus experiments."),
                ("DPDP Act 2023. ", "Because no personal data is collected, processing under the Act is minimal. If user accounts or personal documents are added, consent, purpose limitation and erasure must be implemented.")]: B(t, lead)
P("Learner Signature: _________________________        Date: ________________", after=10)

# ============================ APPENDIX
H1("Appendix")
H2("Appendix A - AI tools disclosure")
T([["Tool", "Where it was used", "Human oversight"],
   ["Claude (Anthropic) chat assistant", "Helped design the architecture, generate the corpus text, Q&A pairs, code, tests, EDA notebook and the draft of this report.", "[Learner to confirm: I ran the code, checked the facts in the corpus against official sources, and edited this report.]"],
   ["LLM APIs / Ollama (run time)", "Used by the finished app to write answers from retrieved chunks.", "Source chunks are shown; the user verifies."],
   ["Matplotlib / seaborn", "Architecture diagrams and EDA figures.", "Checked against the implementation."]], widths=[4, 7, 5.5])
H2("Appendix B - Prompt design")
P("System prompt rules: (1) answer only from the numbered context passages; (2) if the answer is not there, reply exactly \"I could not find this in the documents.\"; (3) keep the answer to 1-4 sentences and cite passage numbers such as [1]; (4) never invent facts, numbers, dates or names; do not give legal or medical advice. Temperature is 0.1. The user message contains the numbered passages, each prefixed with its document title, followed by the question.")
H2("Appendix C - Output format (illustrative)")
P("Question: When did the Constitution of India come into force?\nAnswer: The Constitution came into force on 26 January 1950 [1].\nSources: [1] Preamble and Making of the Constitution | DOC001 (TXT) | score 0.62 - \"The Constitution came into force on 26 January 1950 ...\"\nTop retrieval score: 0.62   Grounding score: 1.00", italic=True)
H2("Appendix D - Installation and run guide")
P("The full, step-by-step guide for Windows, macOS and Linux is in README.md (virtual environment, pip install, building the corpus, running the app, evaluation, notebook, tests and troubleshooting). Short version: python -m venv venv; activate it; pip install -r requirements.txt; python scripts/build_corpus.py; streamlit run app.py.")
H2("Appendix E - Project file structure")
P("rag-doc-qa/\n  app.py  requirements.txt  README.md  .env.example\n  rag/ (loaders, embeddings, store, llm, pipeline, logdb)\n  scripts/ (build_corpus, ingest_corpus, evaluate, ablation, refusal_test, expand_dataset, clean_dataset, make_data_dictionary, ...)\n  data/ (corpus/, dataset.csv, qa_pairs.csv, eval and ablation results)\n  data_dictionary.csv   notebooks/EDA.ipynb   figures/   tests/   docs/   report/", italic=True, size=9.5)
H2("Appendix F - Future enhancements")
for t in ["Hybrid search (keyword BM25 plus vectors) and a cross-encoder re-ranker.", "OCR for scanned PDFs and table-aware parsing.", "Answer evaluation with an LLM judge and RAGAS-style faithfulness metrics.",
          "Conversation memory and follow-up questions.", "Multi-user login, document-level access control and a shared index.", "Support for Hindi and other Indian languages with multilingual embeddings."]: B(t)
H2("Appendix G - Screenshots")
P("[Paste screenshots here: (1) sidebar settings and Documents tab, (2) answer with source chunks, (3) refusal of an out-of-domain question, (4) Analytics tab, (5) Evaluation tab.]", italic=True)
doc.save(os.path.join(ROOT, "report", "Project_Report_RAG_Document_QA.docx")); print("saved")
